"""Knowledge bases: reading each kind of file, chunking, uploading and
processing in the background, search, and cited answers. Embeddings, the
vision model and the answering model are all faked."""

import asyncio
import io
import json
import time

import pytest
from fastapi.testclient import TestClient
from langchain_core.language_models.fake_chat_models import GenericFakeChatModel
from langchain_core.messages import AIMessage, HumanMessage, SystemMessage

from polly_server import embeddings, knowledge, models
from polly_server.api.app import app
from polly_server.api.routers import knowledge as knowledge_router
from polly_server.knowledge import answering, extract, ingest, store
from polly_server.knowledge.chunking import OVERLAP, SIZE, Page, chunk_pages
from tests.test_retrieval import fake_vector

# ---------- fixtures ----------


@pytest.fixture(autouse=True)
def own_data(configure, tmp_path):
    configure(data_dir=tmp_path)


@pytest.fixture(autouse=True)
def fakes(monkeypatch):
    """Embeddings by bag of words; the vision model reads every image as the
    same words."""
    seen = {"images": []}

    async def embed(texts, *, kind):
        return [fake_vector(t) for t in texts]

    async def transcribe(image, mime):
        seen["images"].append((image[:8], mime))
        return "Transcribed: the warranty lasts five years."

    monkeypatch.setattr(embeddings, "embed", embed)
    monkeypatch.setattr(extract, "transcribe", transcribe)
    return seen


_SEEN: list = []


class Reader(GenericFakeChatModel):
    """Streams its reply word by word and keeps the messages it was sent."""

    def _stream(self, messages, *args, **kwargs):
        _SEEN.append(messages)
        yield from super()._stream(messages, *args, **kwargs)


@pytest.fixture
def answers(monkeypatch, configure):
    configure(nebius_api_key="test-key")
    _SEEN.clear()
    made = {}

    def chat_model(*args, **kwargs):
        made.update(kwargs)
        return Reader(messages=iter([AIMessage(content=made["reply"])]))

    monkeypatch.setattr(models, "chat_model", chat_model)

    def reply(text):
        made["reply"] = text
        return made

    return reply


# ---------- files ----------


def make_pdf(pages: list[str]) -> bytes:
    """A small PDF with one Helvetica text line per line; "" is a blank
    page, as a scan with no text layer looks to a text extractor."""
    objects = [b"<< /Type /Catalog /Pages 2 0 R >>"]
    page_ids = [4 + 2 * i for i in range(len(pages))]
    kids = " ".join(f"{p} 0 R" for p in page_ids)
    objects.append(f"<< /Type /Pages /Kids [{kids}] /Count {len(pages)} >>".encode())
    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    for page_id, text in zip(page_ids, pages, strict=True):
        lines = [line.replace("(", "[").replace(")", "]") for line in text.split("\n")]
        ops = ["BT /F1 11 Tf 14 TL 72 740 Td", *(f"({line}) Tj T*" for line in lines), "ET"]
        stream = "\n".join(ops).encode() if text else b""
        objects.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            f"/Resources << /Font << /F1 3 0 R >> >> /Contents {page_id + 1} 0 R >>".encode()
        )
        objects.append(b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream")
    out = b"%PDF-1.4\n"
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{number} 0 obj\n".encode() + body + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    out += b"".join(f"{offset:010d} 00000 n \n".encode() for offset in offsets)
    out += (
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n"
    ).encode()
    return out


def make_docx() -> bytes:
    import docx

    document = docx.Document()
    document.add_heading("Budget 2026", level=1)
    document.add_paragraph("The marketing budget is 42 thousand euros.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Team", "Amount"
    table.cell(1, 0).text, table.cell(1, 1).text = "Design", "12"
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


def make_xlsx() -> bytes:
    from openpyxl import Workbook

    book = Workbook()
    sheet = book.active
    sheet.title = "Q3"
    sheet.append(["Region", "Revenue"])
    sheet.append(["North", 120])
    book.create_sheet("Empty")
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def make_png() -> bytes:
    from PIL import Image

    out = io.BytesIO()
    Image.new("RGB", (40, 20), "white").save(out, "PNG")
    return out.getvalue()


# ---------- reading files ----------


def test_chunks_keep_their_page_and_overlap():
    text = " ".join(f"Sentence number {i} says something." for i in range(120))
    pages = [Page(1, text), Page(2, "Short second page.")]
    chunks = chunk_pages(pages)

    first_page = [c for c in chunks if c.page == 1]
    assert len(first_page) >= 3
    assert all(len(c.text) <= SIZE + OVERLAP for c in chunks)
    assert [c.index for c in chunks] == list(range(len(chunks)))
    # Each chunk starts with the end of the one before, and ends a sentence.
    for before, after in zip(first_page, first_page[1:], strict=False):
        assert after.text[:40] in before.text[-OVERLAP:]
        assert before.text.endswith("says something.")
    assert chunks[-1].page == 2 and chunks[-1].text == "Short second page."

    # A run of text with nowhere to break is still cut.
    (long,) = [Page(None, "x" * 3000)]
    assert [len(c.text) for c in chunk_pages([long])][0] <= SIZE + OVERLAP
    assert chunk_pages([Page(None, "  \n\n ")]) == []


def test_every_kind_of_file_is_read(tmp_path, fakes):
    def read(name, data):
        path = tmp_path / name
        path.write_bytes(data)
        return asyncio.run(extract.extract(path))

    pdf = read("report.pdf", make_pdf(["Revenue grew in the north.\nCosts fell.", ""]))
    assert [p.number for p in pdf] == [1, 2]
    assert "Revenue grew in the north." in pdf[0].text and "Costs fell." in pdf[0].text
    # The blank page was rendered and read as an image.
    assert pdf[1].text.startswith("Transcribed:")
    assert fakes["images"] == [(b"\x89PNG\r\n\x1a\n", "image/png")]

    (image,) = read("photo.png", make_png())
    assert image.number is None and image.text.startswith("Transcribed:")

    (word,) = read("budget.docx", make_docx())
    assert word.text.startswith("# Budget 2026\n\nThe marketing budget is 42 thousand euros.")
    assert "Team | Amount\nDesign | 12" in word.text

    (sheet,) = read("sales.xlsx", make_xlsx())
    assert sheet.text == "## Sheet: Q3\n\nRegion,Revenue\nNorth,120"

    (page,) = read(
        "page.html",
        b"<html><head><title>x</title><script>var a</script></head><body>"
        b"<h1>Return policy</h1><p>Items can be <b>returned</b> in 30 days.</p></body></html>",
    )
    assert page.text == "Return policy\n\nItems can be returned in 30 days."

    (data,) = read("data.json", b'{"a":{"b":[1,2]}}')
    assert json.loads(data.text) == {"a": {"b": [1, 2]}} and "\n" in data.text

    (text,) = read("notes.md", "# Notes\n\nCafé at 9.".encode())
    assert text.text == "# Notes\n\nCafé at 9."
    (latin,) = read("old.txt", "Café".encode("cp1252"))
    assert latin.text == "Café"

    with pytest.raises(extract.Unreadable, match="can't read .xyz files"):
        read("thing.xyz", b"?")


# ---------- the HTTP API ----------


def _wait(client, kb_id, timeout=15):
    deadline = time.time() + timeout
    while time.time() < deadline:
        found = client.get(f"/knowledge/{kb_id}").json()
        if all(f["status"] != "processing" for f in found["files"]):
            return found
        time.sleep(0.05)
    raise AssertionError(f"files still processing: {found}")


def _upload(client, kb_id, *files):
    res = client.post(
        f"/knowledge/{kb_id}/files",
        files=[("files", (name, data, mime)) for name, data, mime in files],
    )
    assert res.status_code == 201, res.text
    return res.json()["files"]


def test_bases_files_search_and_delete():
    with TestClient(app) as client:
        assert client.get("/knowledge").json() == {"bases": []}
        made = client.post("/knowledge", json={"name": " Product ", "description": "Docs"})
        assert made.status_code == 201
        base = made.json()
        assert base["name"] == "Product" and base["files"] == 0 and base["tokens"] == 0
        other = client.post("/knowledge", json={"name": "Other"}).json()
        assert [b["id"] for b in client.get("/knowledge").json()["bases"]] == [
            other["id"],
            base["id"],
        ]

        uploaded = _upload(
            client,
            base["id"],
            ("notes.txt", b"The office wifi password is on the fridge.", "text/plain"),
            ("report.pdf", make_pdf(["Revenue grew in the north.", ""]), "application/pdf"),
            ("budget.docx", make_docx(), "application/octet-stream"),
            ("archive.xyz", b"???", "application/octet-stream"),
        )
        assert [f["status"] for f in uploaded] == ["processing"] * 3 + ["error"]
        assert "can't read .xyz files" in uploaded[3]["error"]
        assert uploaded[0]["mime"] == "text/plain" and uploaded[0]["size"] == 42
        assert uploaded[2]["mime"].endswith("wordprocessingml.document")

        detail = _wait(client, base["id"])
        files = {f["name"]: f for f in detail["files"]}
        assert [f["status"] for f in files.values()].count("ready") == 3
        assert files["report.pdf"]["pages"] == 2 and files["notes.txt"]["pages"] is None
        assert all(f["tokens"] > 0 for n, f in files.items() if n != "archive.xyz")
        assert detail["base"]["files"] == 4
        assert detail["base"]["tokens"] == sum(f["tokens"] for f in files.values())
        assert client.get("/knowledge").json()["bases"][0]["id"] == base["id"]
        assert (store.upload_dir(files["notes.txt"]["id"]) / "notes.txt").exists()

        hits = client.post(
            f"/knowledge/{base['id']}/search", json={"query": "marketing budget"}
        ).json()["hits"]
        assert hits[0]["file_name"] == "budget.docx" and "42 thousand" in hits[0]["text"]
        assert hits[0]["chunk_id"].startswith(files["budget.docx"]["id"])
        scanned = client.post(
            f"/knowledge/{base['id']}/search", json={"query": "warranty", "k": 1}
        ).json()["hits"]
        assert scanned[0]["file_name"] == "report.pdf" and scanned[0]["page"] == 2

        renamed = client.patch(f"/knowledge/{base['id']}", json={"name": "Handbook"}).json()
        assert renamed["name"] == "Handbook" and renamed["description"] == "Docs"
        assert client.patch(f"/knowledge/{base['id']}", json={"name": ""}).status_code == 422

        pdf_id = files["report.pdf"]["id"]
        assert client.delete(f"/knowledge/{base['id']}/files/{pdf_id}").status_code == 204
        assert client.delete(f"/knowledge/{base['id']}/files/{pdf_id}").status_code == 404
        assert store.index().count({"file_id": pdf_id}) == 0
        assert not store.upload_dir(pdf_id).exists() and not store.text_path(pdf_id).exists()
        assert len(client.get(f"/knowledge/{base['id']}").json()["files"]) == 3

        assert client.delete(f"/knowledge/{base['id']}").status_code == 204
        assert client.get(f"/knowledge/{base['id']}").status_code == 404
        assert store.index().count({"kb_id": base["id"]}) == 0
        assert not store.upload_dir(files["notes.txt"]["id"]).exists()
        assert client.delete(f"/knowledge/{base['id']}").status_code == 404
        assert client.post("/knowledge/nope/search", json={"query": "x"}).status_code == 404


def test_uploads_over_the_limit_are_refused(monkeypatch):
    monkeypatch.setattr(knowledge_router, "MAX_UPLOAD", 10)
    with TestClient(app) as client:
        base = client.post("/knowledge", json={"name": "Small"}).json()
        res = client.post(
            f"/knowledge/{base['id']}/files",
            files=[("files", ("big.txt", b"x" * 20, "text/plain"))],
        )
        assert res.status_code == 413
        assert client.get(f"/knowledge/{base['id']}").json()["files"] == []
        assert not list((store.root() / "files").glob("*/*"))


def test_a_file_left_processing_is_picked_up_again():
    base = store.create_base("Restarted")
    file = store.new_file(base.id, "left.txt", "text/plain")
    store.upload_path(file).parent.mkdir(parents=True)
    store.upload_path(file).write_bytes(b"Left behind when the server stopped.")
    store.add_file(file)
    with TestClient(app) as client:
        detail = _wait(client, base.id)
    assert detail["files"][0]["status"] == "ready"


def test_unreadable_files_say_why():
    with TestClient(app) as client:
        base = client.post("/knowledge", json={"name": "Broken"}).json()
        _upload(
            client,
            base["id"],
            ("empty.txt", b"   ", "text/plain"),
            ("broken.pdf", b"%PDF-1.4 not really", "application/pdf"),
        )
        files = {f["name"]: f for f in _wait(client, base["id"])["files"]}
    assert files["empty.txt"]["error"] == "Polly found no text in this file."
    assert files["broken.pdf"]["status"] == "error" and files["broken.pdf"]["error"]


# ---------- answers ----------


def _frames(body: str) -> list[dict]:
    return [
        json.loads(line[len("data: ") :]) for line in body.splitlines() if line.startswith("data: ")
    ]


def test_chat_streams_cited_sources_then_the_answer(answers):
    answers("The wifi password is on the fridge [1].")
    with TestClient(app) as client:
        base = client.post("/knowledge", json={"name": "Home"}).json()
        _upload(
            client,
            base["id"],
            ("notes.txt", b"The office wifi password is on the fridge.", "text/plain"),
            ("menu.txt", b"Soup on Mondays.", "text/plain"),
        )
        _wait(client, base["id"])
        res = client.post(
            f"/knowledge/{base['id']}/chat",
            json={
                "message": "Where is the wifi password?",
                "history": [
                    {"role": "user", "text": "Hi"},
                    {"role": "assistant", "text": "Hello."},
                ],
            },
        )
    assert res.status_code == 200 and res.headers["content-type"].startswith("text/event-stream")
    frames = _frames(res.text)
    kinds = [f["type"] for f in frames]
    assert kinds[0] == "sources" and kinds[-2:] == ["sources", "done"]
    assert set(kinds[1:-2]) == {"delta"} and len(kinds) > 4

    every = frames[0]["sources"]
    assert [(s["n"], s["file_name"]) for s in every] == [(1, "notes.txt"), (2, "menu.txt")]
    assert every[0]["excerpt"] == "The office wifi password is on the fridge."
    assert [s["n"] for s in frames[-2]["sources"]] == [1]
    assert frames[-1]["text"] == "The wifi password is on the fridge [1]."
    assert "".join(f["text"] for f in frames if f["type"] == "delta") == frames[-1]["text"]

    # The whole base fits: the model read every passage, then the history.
    (messages,) = _SEEN
    assert isinstance(messages[0], SystemMessage)
    assert answering.WHOLE in messages[0].content
    assert '<passage n="2" file="menu.txt">\nSoup on Mondays.\n</passage>' in messages[0].content
    assert [m.content for m in messages[1:]] == ["Hi", "Hello.", "Where is the wifi password?"]


def test_big_bases_are_searched_and_the_screen_is_context(answers, monkeypatch):
    made = answers("Five years [1][2]; see also [9].")
    monkeypatch.setattr(answering, "LONG_CONTEXT_TOKENS", 0)

    async def go():
        base = store.create_base("Big")
        for name, text in [
            ("warranty.txt", "The warranty lasts five years from delivery."),
            ("menu.txt", "Soup on Mondays."),
        ]:
            file = store.new_file(base.id, name, "text/plain")
            store.upload_path(file).parent.mkdir(parents=True)
            store.upload_path(file).write_text(text)
            store.add_file(file)
            ingest.schedule(file)
        await ingest.wait()
        return base, [
            e
            async for e in knowledge.answer(
                [base.id, "kb-missing"],
                "How long is the warranty?",
                [],
                context="Email from Maria: is it still under warranty?",
            )
        ]

    base, events = asyncio.run(go())
    assert made["model"] == answering.settings.strong_model
    assert made["reasoning_effort"] == "medium"
    first = events[0]["sources"]
    assert first[0]["file_name"] == "warranty.txt"
    assert [s["n"] for s in events[-2]["sources"]] == [1, 2]
    (messages,) = _SEEN
    assert answering.SOME in messages[0].content
    assert isinstance(messages[-1], HumanMessage)
    assert messages[-1].content.startswith("What is on my screen right now:\n<screen>\nEmail from")
    assert messages[-1].content.endswith("How long is the warranty?")

    hits = asyncio.run(knowledge.search([base.id], "warranty", k=1))
    assert hits[0].file_name == "warranty.txt" and hits[0].page is None
    assert [b.name for b in knowledge.list_bases()] == ["Big"]


def test_answers_need_files_and_a_key(answers, configure):
    base = store.create_base("Empty")

    async def collect(kb_ids):
        return [e async for e in knowledge.answer(kb_ids, "Anything?", [])]

    events = asyncio.run(collect([base.id]))
    assert events[0] == {"type": "sources", "sources": []}
    assert events[1]["type"] == "done" and "no readable files" in events[1]["text"]
    assert _SEEN == []
    assert asyncio.run(collect(["kb-nope"]))[0]["type"] == "error"

    with TestClient(app) as client:
        assert client.post("/knowledge/kb-nope/chat", json={"message": "x"}).status_code == 404
        configure(nebius_api_key="")
        res = client.post(f"/knowledge/{base.id}/chat", json={"message": "x"})
        assert res.status_code == 503


def test_a_failing_model_ends_the_stream_with_an_error(monkeypatch, configure):
    configure(nebius_api_key="test-key")

    def broken(*args, **kwargs):
        raise RuntimeError("Token Factory is down")

    monkeypatch.setattr(models, "chat_model", broken)
    base = store.create_base("Docs")
    file = store.new_file(base.id, "a.txt", "text/plain")
    store.add_file(file)
    store.update_file(file.id, status="ready", tokens=3)
    store.save_text(file.id, [Page(None, "Some text.")])

    async def collect():
        return [e async for e in knowledge.answer([base.id], "Question?", [])]

    events = asyncio.run(collect())
    assert [e["type"] for e in events] == ["sources", "error"]
    assert "Token Factory is down" in events[-1]["message"]


def test_citations_are_read_in_every_form():
    text = "See [2] and [1, 3], then [4-5] or [4–5]; not [99] or [0] or [x]."
    assert answering.cited(text, 5) == [2, 1, 3, 4, 5]
