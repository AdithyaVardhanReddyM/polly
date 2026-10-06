"""Reading an uploaded file as text, page by page.

PDFs give their text per page (pypdf). A page with next to no text is a scan
or a picture, so it is rendered (pdfium) and read by the vision model, as are
image files. Word, Excel, HTML and plain-text formats are read directly.
Excel sheets come out as CSV-like rows, one section per sheet.
"""

from __future__ import annotations

import asyncio
import base64
import csv
import io
import json
import logging
import re
import threading
from pathlib import Path

from langchain_core.messages import HumanMessage

from polly_server.config import settings
from polly_server.knowledge.chunking import Page

log = logging.getLogger(__name__)

# A PDF page with fewer characters than this is read as an image.
SCANNED_BELOW = 25
# Scanned pages read at once.
VISION_CONCURRENCY = 4
# Images go to the vision model at most this many pixels on a side.
MAX_SIDE = 2_048

KINDS: dict[str, str] = {
    ".pdf": "pdf",
    ".docx": "docx",
    ".xlsx": "xlsx",
    ".xlsm": "xlsx",
    ".png": "image",
    ".jpg": "image",
    ".jpeg": "image",
    ".webp": "image",
    ".gif": "image",
    ".html": "html",
    ".htm": "html",
    ".txt": "text",
    ".md": "text",
    ".markdown": "text",
    ".csv": "text",
    ".tsv": "text",
    ".json": "json",
}
IMAGE_TYPES = {
    ".png": "image/png",
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
    ".webp": "image/webp",
    ".gif": "image/gif",
}

TRANSCRIBE_PROMPT = """\
Transcribe all the text in this image, in reading order, as Markdown: keep \
headings, lists and tables. Describe each figure, chart or photo in one line, \
in italics. Reply with the transcription only."""

# pdfium is not thread-safe, even across documents.
_pdfium = threading.Lock()


class Unreadable(ValueError):
    """The file cannot be read; the message says why, for the user."""


def kind_of(name: str) -> str | None:
    return KINDS.get(Path(name).suffix.casefold())


def unsupported_message(name: str) -> str:
    suffix = Path(name).suffix or "these"
    return (
        f"Polly can't read {suffix} files. Add PDF, Word (.docx), Excel (.xlsx), text, "
        "Markdown, CSV, JSON, HTML or image files."
    )


async def extract(path: Path) -> list[Page]:
    """The file's text as pages. Raises `Unreadable`."""
    kind = kind_of(path.name)
    if kind == "pdf":
        return await _pdf(path)
    if kind == "image":
        mime = IMAGE_TYPES[path.suffix.casefold()]
        return [Page(None, await transcribe(path.read_bytes(), mime))]
    if kind == "docx":
        return [Page(None, await asyncio.to_thread(_docx, path))]
    if kind == "xlsx":
        return [Page(None, await asyncio.to_thread(_xlsx, path))]
    if kind in {"text", "json", "html"}:
        return [Page(None, await asyncio.to_thread(_plain, path, kind))]
    raise Unreadable(unsupported_message(path.name))


# ---------- the vision model ----------


def _for_vision(data: bytes, mime: str) -> tuple[bytes, str]:
    """Shrink big images and turn GIFs into PNG, which every endpoint takes."""
    from PIL import Image

    with Image.open(io.BytesIO(data)) as image:
        if max(image.size) <= MAX_SIDE and mime != "image/gif":
            return data, mime
        image.seek(0)
        frame = image.convert("RGBA" if image.mode in {"RGBA", "LA", "P"} else "RGB")
        frame.thumbnail((MAX_SIDE, MAX_SIDE))
        out = io.BytesIO()
        frame.save(out, "PNG")
        return out.getvalue(), "image/png"


def _text_of(content: object) -> str:
    if isinstance(content, str):
        return content
    if isinstance(content, list):
        return "".join(
            b.get("text", "") if isinstance(b, dict) else str(b)
            for b in content
            if not isinstance(b, dict) or b.get("type", "text") == "text"
        )
    return ""


async def transcribe(image: bytes, mime: str) -> str:
    """What the vision model reads in an image, as Markdown."""
    from polly_server.models import chat_model

    data, mime = await asyncio.to_thread(_for_vision, image, mime)
    model = chat_model(model=settings.vision_model, reasoning_effort="low")
    url = f"data:{mime};base64,{base64.b64encode(data).decode()}"
    reply = await model.ainvoke(
        [
            HumanMessage(
                content=[
                    {"type": "text", "text": TRANSCRIBE_PROMPT},
                    {"type": "image_url", "image_url": {"url": url}},
                ]
            )
        ]
    )
    text = _text_of(reply.content).strip()
    # Some replies come wrapped in a Markdown fence.
    fenced = re.fullmatch(r"```(?:markdown|md)?\s*\n(.*)\n```", text, re.DOTALL)
    return (fenced.group(1) if fenced else text).strip()


# ---------- PDF ----------


def _pdf_text(path: Path) -> list[str]:
    from pypdf import PdfReader
    from pypdf.errors import DependencyError, PdfReadError

    try:
        reader = PdfReader(path)
        if reader.is_encrypted and not reader.decrypt(""):
            raise Unreadable("This PDF has a password; remove it and add the file again.")
        pages = list(reader.pages)
    except (PdfReadError, DependencyError) as exc:
        raise Unreadable(f"This PDF could not be opened ({exc}).") from exc
    texts: list[str] = []
    for page in pages:
        try:
            texts.append(page.extract_text() or "")
        except Exception as exc:  # noqa: BLE001 - one bad page should not lose the rest
            log.warning("%s: a page could not be read: %s", path.name, exc)
            texts.append("")
    return texts


def _render(path: Path, index: int) -> bytes:
    """One PDF page as a PNG, at most `MAX_SIDE` pixels on a side."""
    import pypdfium2 as pdfium

    with _pdfium:
        pdf = pdfium.PdfDocument(str(path))
        try:
            page = pdf[index]
            width, height = page.get_size()
            scale = min(2.0, MAX_SIDE / max(width, height, 1))
            image = page.render(scale=scale).to_pil()
        finally:
            pdf.close()
    out = io.BytesIO()
    image.save(out, "PNG")
    return out.getvalue()


async def _pdf(path: Path) -> list[Page]:
    texts = await asyncio.to_thread(_pdf_text, path)
    scanned = [i for i, text in enumerate(texts) if len(text.strip()) < SCANNED_BELOW]
    gate = asyncio.Semaphore(VISION_CONCURRENCY)

    async def read(i: int) -> None:
        async with gate:
            try:
                png = await asyncio.to_thread(_render, path, i)
                texts[i] = await transcribe(png, "image/png")
            except Exception as exc:  # noqa: BLE001 - keep what the page had
                log.warning("%s: page %d could not be transcribed: %s", path.name, i + 1, exc)

    await asyncio.gather(*(read(i) for i in scanned))
    return [Page(i + 1, text) for i, text in enumerate(texts)]


# ---------- office files ----------


def _docx(path: Path) -> str:
    import docx
    from docx.table import Table

    document = docx.Document(str(path))
    blocks: list[str] = []
    for block in document.iter_inner_content():
        if isinstance(block, Table):
            rows = []
            for row in block.rows:
                cells: list[str] = []
                for cell in row.cells:
                    text = " ".join(cell.text.split())
                    # A merged cell repeats across the columns it spans.
                    if not cells or cells[-1] != text:
                        cells.append(text)
                rows.append(" | ".join(cells))
            blocks.append("\n".join(rows))
            continue
        text = block.text.strip()
        if not text:
            continue
        style = block.style.name if block.style is not None else ""
        level = re.fullmatch(r"Heading (\d)", style or "")
        if level:
            text = f"{'#' * int(level.group(1))} {text}"
        elif style == "Title":
            text = f"# {text}"
        elif "List" in (style or ""):
            text = f"- {text}"
        blocks.append(text)
    return "\n\n".join(blocks)


def _xlsx(path: Path) -> str:
    from openpyxl import load_workbook

    book = load_workbook(path, read_only=True, data_only=True)
    sections: list[str] = []
    try:
        for sheet in book.worksheets:
            out = io.StringIO()
            writer = csv.writer(out, lineterminator="\n")
            for row in sheet.iter_rows(values_only=True):
                values = list(row)
                while values and values[-1] is None:
                    values.pop()
                if values:
                    writer.writerow(["" if v is None else v for v in values])
            if rows := out.getvalue().strip():
                sections.append(f"## Sheet: {sheet.title}\n\n{rows}")
    finally:
        book.close()
    return "\n\n".join(sections)


# ---------- text ----------


def _decode(data: bytes) -> str:
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        return data.decode("cp1252", errors="replace")


def _plain(path: Path, kind: str) -> str:
    text = _decode(path.read_bytes())
    if kind == "json":
        try:
            # Minified JSON has no line breaks to cut chunks at.
            return json.dumps(json.loads(text), indent=2, ensure_ascii=False)
        except ValueError:
            return text
    if kind == "html":
        from bs4 import BeautifulSoup

        soup = BeautifulSoup(text, "html.parser")
        for tag in soup(["script", "style", "noscript", "template", "svg", "head"]):
            tag.decompose()
        # Blocks become paragraphs; inline tags stay inside their sentence.
        for tag in soup.find_all(_HTML_BLOCKS):
            tag.insert_before("\n\n")
            tag.insert_after("\n\n")
        for tag in soup.find_all(["br", "tr"]):
            tag.insert_after("\n")
        for tag in soup.find_all(["td", "th"]):
            tag.insert_after(" | ")
        lines = (" ".join(line.split()) for line in soup.get_text().splitlines())
        return re.sub(r"\n{3,}", "\n\n", "\n".join(lines)).strip()
    return text


_HTML_BLOCKS = [
    "p",
    "div",
    "section",
    "article",
    "header",
    "footer",
    "li",
    "h1",
    "h2",
    "h3",
    "h4",
    "h5",
    "h6",
    "pre",
    "blockquote",
    "table",
    "ul",
    "ol",
    "dl",
    "dt",
    "dd",
    "figure",
]
